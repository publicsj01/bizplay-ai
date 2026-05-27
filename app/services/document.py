# 문서 업로드 → 텍스트 추출 → 청킹 → 임베딩 저장 파이프라인
# Spring AI: DocumentService.java + DocumentParserService.java 1:1 대응

import asyncio
import io
import re
import uuid
from pathlib import Path
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func

from app.core.config import settings
from app.models.entities import Bot, Document
from app.services.vector import embed_and_store, delete_by_doc_id


# ── 텍스트 추출 ────────────────────────────────────────────────────────────────

def extract_text_from_pdf(content: bytes) -> str:
    """Spring AI: DocumentParserService.parsePdf(PDFBox) 대응"""
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(content: bytes) -> str:
    """Spring AI: DocumentParserService.parseDocx(Apache POI) 대응"""
    from docx import Document as DocxDocument
    doc = DocxDocument(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text(content: bytes, content_type: str) -> str:
    """MIME 타입별 텍스트 추출 디스패처"""
    if "pdf" in content_type:
        text = extract_text_from_pdf(content)
    elif "word" in content_type or "docx" in content_type or "openxmlformats" in content_type:
        text = extract_text_from_docx(content)
    else:
        text = content.decode("utf-8", errors="replace")

    # Spring AI: NUL 바이트 제거 (PostgreSQL TEXT/JSONB에서 0x00 거부)
    return text.replace("\x00", "")


# ── 청킹 ──────────────────────────────────────────────────────────────────────

def chunk_text(
    text: str,
    chunk_size: int = None,
    chunk_overlap: int = None,
) -> list[str]:
    """
    단락 경계 우선 청킹 + overlap.
    Spring AI: DocumentService.chunkText() 대응
    마크다운 테이블은 헤더를 각 청크에 반복 삽입 (Spring AI 동일 로직)
    """
    size = chunk_size or settings.rag_chunk_size
    overlap = chunk_overlap or settings.rag_chunk_overlap

    # 단락 분리
    paragraphs = re.split(r"\n\s*\n", text.strip())
    chunks: list[str] = []
    current = ""

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue

        # 마크다운 테이블 처리: 행 단위 분리 후 헤더 반복
        if "|" in para and re.match(r"^\|.+\|", para):
            rows = para.splitlines()
            header = rows[0] if rows else ""
            sep = rows[1] if len(rows) > 1 else ""
            table_chunk = ""
            for row in rows[2:] if len(rows) > 2 else []:
                candidate = f"{header}\n{sep}\n{table_chunk}{row}\n"
                if len(candidate) > size and table_chunk:
                    chunks.append(f"{header}\n{sep}\n{table_chunk}".strip())
                    table_chunk = row + "\n"
                else:
                    table_chunk += row + "\n"
            if table_chunk:
                chunks.append(f"{header}\n{sep}\n{table_chunk}".strip())
            continue

        if len(current) + len(para) + 2 > size:
            if current:
                chunks.append(current.strip())
                # overlap: 이전 청크 끝부분을 다음 청크 시작에 포함
                current = current[-overlap:] + "\n\n" + para if overlap > 0 else para
            else:
                # 단락 자체가 chunk_size 초과 → 강제 분할
                for i in range(0, len(para), size - overlap):
                    chunks.append(para[i : i + size])
                current = ""
        else:
            current = (current + "\n\n" + para).strip() if current else para

    if current.strip():
        chunks.append(current.strip())

    return [c for c in chunks if c.strip()]


# ── 메인 파이프라인 ───────────────────────────────────────────────────────────

async def ingest_document(
    db: AsyncSession,
    bot_id: uuid.UUID,
    file_content: bytes,
    file_name: str,
    content_type: str,
    title: str,
) -> Document:
    """
    문서 업로드 전체 파이프라인.
    Spring AI: DocumentService.upload() 대응
    1. Document 레코드 생성 (PROCESSING)
    2. 파일 디스크 저장
    3. 텍스트 추출
    4. 청킹
    5. 임베딩 → PGVector 저장
    6. 상태 COMPLETED/FAILED 업데이트
    """
    doc = Document(
        bot_id=bot_id,
        title=title,
        file_name=file_name,
        content_type=content_type,
        embedding_status="PROCESSING",
    )
    db.add(doc)
    await db.flush()  # doc.id 확보

    # 파일 저장: uploads/{doc_id}/{originalFileName}
    upload_path = Path(settings.upload_dir) / str(doc.id)
    upload_path.mkdir(parents=True, exist_ok=True)
    file_path = upload_path / file_name
    file_path.write_bytes(file_content)

    # 백그라운드에서 임베딩 처리 (상태 업데이트는 별도 세션에서)
    doc_id = doc.id
    asyncio.create_task(
        _embed_in_background(doc_id, bot_id, file_content, content_type, title, file_name)
    )

    return doc


async def _embed_in_background(
    doc_id: uuid.UUID,
    bot_id: uuid.UUID,
    file_content: bytes,
    content_type: str,
    title: str,
    file_name: str,
) -> None:
    """임베딩 작업을 백그라운드에서 실행하고 Document 상태를 업데이트한다."""
    from app.core.database import AsyncSessionLocal

    async with AsyncSessionLocal() as session:
        try:
            text = extract_text(file_content, content_type)
            chunks = chunk_text(text)

            metadata_list: list[dict[str, Any]] = [
                {
                    "bot_id": str(bot_id),
                    "doc_id": str(doc_id),
                    "title": title,
                    "file_name": file_name,
                    "content_type": content_type,
                    "chunk_index": idx,
                }
                for idx, _ in enumerate(chunks)
            ]

            await embed_and_store(chunks, metadata_list)

            result = await session.get(Document, doc_id)
            if result:
                result.embedding_status = "COMPLETED"
                await session.commit()

        except Exception:
            result = await session.get(Document, doc_id)
            if result:
                result.embedding_status = "FAILED"
                await session.commit()


async def delete_document(db: AsyncSession, doc_id: uuid.UUID) -> None:
    """
    문서 레코드 + 벡터 청크 + 디스크 파일 삭제.
    Spring AI: DocumentService.delete() 대응
    """
    doc = await db.get(Document, doc_id)
    if not doc:
        return

    # 벡터 청크 삭제
    await delete_by_doc_id(str(doc_id))

    # 디스크 파일 삭제
    file_path = Path(settings.upload_dir) / str(doc_id) / doc.file_name
    if file_path.exists():
        file_path.unlink()
    parent = file_path.parent
    if parent.exists() and not any(parent.iterdir()):
        parent.rmdir()

    await db.delete(doc)
